# -*- coding: utf-8 -*-
"""Generate the concise highlights brief (AI Trading Automation 2.1.3) as a styled .docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

DST = r"C:\Users\zheng siyuan\Desktop\AI Trading Automation 2.1.3 核心亮点简报.docx"

TEAL = "10272B"        # product dark teal
TEAL_RGB = RGBColor(0x10, 0x27, 0x2B)
ACCENT = RGBColor(0x49, 0xB9, 0xA5)   # product highlight green
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
FILL_LIGHT = "EFF6F4"   # light mint card fill
FILL_CARD = "F5F7FA"
FILL_TABLE_H = "D9E2F3"

BODY_EA = "微软雅黑"


def set_ea(run, ea=BODY_EA):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), ea)


def shade_par(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_border(cell, **kwargs):
    """kwargs: top/bottom/left/right = (val, sz, color)"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, (val, sz, color) in kwargs.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)


def heading_bar(doc, text, num):
    """Section heading rendered as a filled teal bar with white bold text."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    shade_par(p, TEAL)
    r = p.add_run(f"  {num}  {text}")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = WHITE
    set_ea(r)
    return p


def card(doc, title, desc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Inches(0.06)
    shade_par(p, FILL_CARD)
    r1 = p.add_run("▍" + title + "  ")
    r1.bold = True
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = TEAL_RGB
    set_ea(r1)
    r2 = p.add_run(desc)
    r2.font.size = Pt(10)
    set_ea(r2)
    return p


def bullet(doc, lead, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("● ").font.color.rgb = ACCENT
    r1 = p.add_run(lead)
    r1.bold = True
    r1.font.color.rgb = TEAL_RGB
    set_ea(r1)
    r2 = p.add_run(text)
    set_ea(r2)
    return p


def body(doc, text, size=10, color=None, bold=False, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    set_ea(r)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_EA)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    # ── title block ────────────────────────────────────────────────────────
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_after = Pt(2)
    shade_par(tp, TEAL)
    r = tp.add_run("\nAI Trading Automation 2.1.3")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = WHITE
    set_ea(r)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(6)
    shade_par(p2, TEAL)
    r = p2.add_run("核心特点与优势速览 · 投资研究与模拟交易桌面应用\n")
    r.font.size = Pt(12)
    r.font.color.rgb = ACCENT
    set_ea(r)

    # one-liner
    onep = doc.add_paragraph()
    onep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    onep.paragraph_format.space_after = Pt(8)
    shade_par(onep, FILL_LIGHT)
    r = onep.add_run("一句话定位：")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = TEAL_RGB
    set_ea(r)
    r = onep.add_run("一个覆盖 A股 / 港股 / 美股 / 场内 ETF 四市场、会自主完成「研究 → 辩论 → 决策 → 风控 → 模拟成交」全流程，且绝不重复成交、绝不碰真实资金的 AI 投资智能体。")
    r.font.size = Pt(11)
    set_ea(r)

    # stats strip
    stats = doc.add_table(rows=1, cols=4)
    stats.alignment = 1
    data = [
        ("4 大市场", "A股 · 港股 · 美股 · ETF"),
        ("13 角色 · 5 阶段", "委员会式分析流水线"),
        ("226 项自动化测试", "Python 184 + Node 22 + C# 20"),
        ("76,167 文件零丢失", "覆盖升级实测验证"),
    ]
    for j, (big, small) in enumerate(data):
        cell = stats.rows[0].cells[j]
        shade_cell(cell, FILL_CARD)
        cell_border(cell, top=("single", 12, TEAL), bottom=("single", 12, TEAL),
                    left=("single", 6, "D0D7E2"), right=("single", 6, "D0D7E2"))
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(big)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = TEAL_RGB
        set_ea(r)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(2)
        r = p2.add_run(small)
        r.font.size = Pt(8.5)
        r.font.color.rgb = GRAY
        set_ea(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # ── 1. 核心特点 ────────────────────────────────────────────────────────
    heading_bar(doc, "六大核心特点", "01")
    card(doc, "四市场一站式，规则内建",
         "A股 / 港股 / 美股 / 场内 ETF 统一行情与选股；T+1/T+0、整手、涨跌停、印花税、佣金、滑点全部内置为确定性规则，模拟撮合按真实市场规则执行。")
    card(doc, "13 角色委员会式投研流程",
         "每只标的并行「技术面/基本面/新闻/情绪」四路研究 → 多空辩论 → 研究经理与交易员 → 组合草案 → 三方风险辩论 → 最终决策。进度、检查点、证据数在「分析流程」页实时可见，崩溃后同轮次从检查点续跑。")
    card(doc, "AI 决策 ≠ 直接执行（核心安全设计）",
         "所有交易必须经过唯一入口，由独立 Python 引擎完成：决策指纹校验 → 硬风控重算仓位 → 纸面模拟撮合。仓位封顶、回撤熔断、硬止损/移动止损/分批止盈优先于 AI 建议——AI 永远绕不过风控代码。")
    card(doc, "永不重复成交",
         "决策指纹 + 账户内执行回执 + 跨进程租约三层幂等：崩溃、超时、重复点击都不会重复成交；失败的完整分析可用同一 ID 从中断处恢复（已实测验证）。")
    card(doc, "密钥与隐私安全",
         "API Key 与 Webhook 用 Windows DPAPI 加密存储，绝不落明文；服务仅监听本机随机端口 + 每次启动随机令牌；数据全本地、开源（MIT），只做模拟交易、不连任何券商。")
    card(doc, "一键安装的桌面产品，还会自我维护",
         "安装包内置 Python/Node/.NET/WebView2，双击即用；单实例、托盘、开机自启、覆盖升级不丢数据。内置自维护能力：AI 可读日志、改源码、跑测试、失败回滚——自己修自己。")

    # ── 2. 全流程一览 ──────────────────────────────────────────────────────
    heading_bar(doc, "完整决策链一览", "02")
    flowp = doc.add_paragraph()
    flowp.paragraph_format.space_after = Pt(6)
    shade_par(flowp, FILL_LIGHT)
    r = flowp.add_run("选股请求 → 硬筛选+六因子评分 → Agent 复筛　|　用户点名股票 → 身份确认（跳过选股）")
    r.font.size = Pt(9.5)
    r.bold = True
    r.font.color.rgb = TEAL_RGB
    set_ea(r)
    flowp.add_run("\n四路基础研究 → 多空辩论 → 研究经理/交易员 → 组合草案 → 三方风险辩论 → 最终决策\n→ 用户批准 → Python 硬风控 → 模拟撮合 → 审计 / 报告 / 反思")
    r = flowp.runs[1]
    r.font.size = Pt(9.5)
    set_ea(r)

    # ── 3. 与同类项目的差异 ────────────────────────────────────────────────
    heading_bar(doc, "与同类项目相比，强在哪里", "03")
    cmp = doc.add_table(rows=6, cols=2)
    cmp.style = doc.styles["Table Grid"]
    rows = [
        ("AI Hedge Fund（LangGraph 多角色）", "止步于美股单市场「给信号」→ IA 四市场 + 完整风控撮合闭环"),
        ("TradingAgents（多智能体投研）", "流程最相似，但止步信号/回测 → IA 是其「工程落地强化版」"),
        ("Qlib / RD-Agent（微软量化平台）", "量化因子与回测更深 → IA 强在 LLM 投研决策与全流程闭环，两者互补"),
        ("OpenBB / FinRobot", "通用数据终端 / 研究工具集 → IA 是端到端决策应用"),
        ("同花顺问财 / 聚宽 / 掘金等", "商业闭源、重实盘导流 → IA 开源、离线、只做模拟"),
        ("底座路线", "同类多自建 LangGraph/AutoGen → IA 复用 DeepSeek Harness 插件化底座，天然获得调度、检查点、插件扩展能力"),
    ]
    widths = (Inches(2.2), Inches(4.9))
    for ri, (a, b) in enumerate(rows):
        c0 = cmp.rows[ri].cells[0]
        c1 = cmp.rows[ri].cells[1]
        c0.width, c1.width = widths
        shade_cell(c0, FILL_CARD)
        p = c0.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
        r = p.add_run(a); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = TEAL_RGB; set_ea(r)
        p = c1.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
        r = p.add_run(b); r.font.size = Pt(9); set_ea(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # ── 4. 为什么可信赖 ────────────────────────────────────────────────────
    heading_bar(doc, "为什么可信赖", "04")
    bullet(doc, "只做纸面：", "引擎代码只接受模拟模式，不连接任何真实券商。")
    bullet(doc, "唯一执行入口：", "AI 只能提交决策，价格由引擎自行抓取、仓位由硬风控重算封顶。")
    bullet(doc, "三层幂等：", "幂等键必填 + 决策指纹内容绑定 + 账户内回执原子写入——重试只会重放结果。")
    bullet(doc, "崩溃可恢复：", "阶段级检查点 + 跨进程租约 + 启动扫描续跑，重启后自动接着做。")
    bullet(doc, "密钥加密：", "DPAPI 加密存储，配置与日志双重脱敏，令牌只注入本机引擎接口。")
    bullet(doc, "质量门禁：", "一键发行门禁（Python/Node/C# 全量测试 + 升级保数据校验），任一失败即拦截发布。")

    # ── 5. 适用场景与边界 ──────────────────────────────────────────────────
    heading_bar(doc, "适用场景与产品边界", "05")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    r = p.add_run("适合："); r.bold = True; r.font.color.rgb = TEAL_RGB; set_ea(r)
    p.add_run("个人投资者的投研学习与模拟盘训练 · 策略验证 · AI 多智能体工程参考 · 开源本地化的投研工具需求。")
    p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(2)
    r = p2.add_run("边界："); r.bold = True; r.font.color.rgb = TEAL_RGB; set_ea(r)
    p2.add_run("仅研究与模拟交易，不构成投资建议，不连接真实券商（产品定位使然）。")

    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(10)
    r = fp.add_run("— 完 · 详细版见《AI Trading Automation 2.1.3 项目报告.docx》 —")
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY
    set_ea(r)

    doc.save(DST)
    print("saved:", DST)


if __name__ == "__main__":
    main()
